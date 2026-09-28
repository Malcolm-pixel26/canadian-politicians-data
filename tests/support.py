"""
Test helpers: a fake network (urllib.request.urlopen), a fake clock
(time.sleep / time.monotonic) so throttling and backoff are measurable
and instant, and routes that serve the Represent / OpenParliament /
House of Commons fixtures (trimmed copies of live responses) and the
Senate / Legislative Assembly of Nunavut fixtures (rebuilt from the live
pages' content -- see the note at the top of each).
"""
import contextlib
import io
import json
import math
import os
import shutil
import sys
import tempfile
import unittest
import urllib.error
import urllib.parse
from unittest import mock
from xml.sax.saxutils import escape

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)
FIXTURES = os.path.join(ROOT, "tests", "fixtures")

import represent_common as rc  # noqa: E402

REPRESENT = "https://represent.opennorth.ca"
OPENPARLIAMENT = "https://api.openparliament.ca"
SENATE = "https://sencanada.ca/umbraco/surface/SenatorsAjax/GetSenators"
NUNAVUT = "https://www.assembly.nu.ca"


def load_fixture(name):
    with open(os.path.join(FIXTURES, name), encoding="utf-8") as f:
        return json.load(f) if name.endswith(".json") else f.read()


class FakeResponse(io.BytesIO):
    def __init__(self, body, status=200):
        super().__init__(body if isinstance(body, bytes) else body.encode("utf-8"))
        self.status = status

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        self.close()


def json_response(payload):
    return FakeResponse(json.dumps(payload))


def http_error(url, code, headers=None):
    return urllib.error.HTTPError(url, code, f"HTTP {code}", headers or {}, None)


class FakeWeb:
    """
    Context manager that replaces urlopen with a router and time with a
    fake clock. Each request advances the clock by `latency` seconds and
    is recorded in `requests` as (time, url). Unrouted URLs get a 404.
    """

    def __init__(self, latency=0.2):
        self.now = 1_000.0
        self.latency = latency
        self.routes = []
        self.requests = []

    # -- clock --
    def monotonic(self):
        return self.now

    def sleep(self, seconds):
        self.now += max(0.0, seconds)

    # -- routing --
    def route(self, match, handler):
        """`match`: URL prefix or predicate; `handler(url)` returns a
        response, or an exception instance to raise."""
        self.routes.insert(0, (match, handler))
        return self

    def urlopen(self, req, timeout=None):
        url = req.full_url if hasattr(req, "full_url") else req
        self.now += self.latency
        self.requests.append((self.now, url))
        for match, handler in self.routes:
            if match(url) if callable(match) else url.startswith(match):
                result = handler(url)
                if isinstance(result, BaseException):
                    raise result
                return result
        raise http_error(url, 404)

    def count(self, fragment):
        return sum(1 for _, url in self.requests if fragment in url)

    def peak_per_window(self, host, window=60.0):
        times = [t for t, url in self.requests if urllib.parse.urlsplit(url).hostname == host]
        return max((sum(1 for t in times if start <= t < start + window) for start in times), default=0)

    def __enter__(self):
        self._patches = [
            mock.patch("urllib.request.urlopen", self.urlopen),
            mock.patch("time.sleep", self.sleep),
            mock.patch("time.monotonic", self.monotonic),
        ]
        for p in self._patches:
            p.start()
        rc.clear_cache()
        rc._next_request_at.clear()
        return self

    def __exit__(self, *exc):
        for p in reversed(self._patches):
            p.stop()
        rc.clear_cache()
        rc._next_request_at.clear()


def _page(objects, url, meta_style="represent"):
    query = urllib.parse.parse_qs(urllib.parse.urlsplit(url).query)
    limit = int(query.get("limit", ["20"])[0])
    offset = int(query.get("offset", ["0"])[0])
    page = objects[offset:offset + limit]
    base = url.split("?", 1)[0]
    more = offset + limit < len(objects)
    next_link = f"{urllib.parse.urlsplit(base).path}?limit={limit}&offset={offset + limit}" if more else None
    return {"objects": page, "meta": {"offset": offset, "limit": limit, "total_count": len(objects),
                                      "previous": None, "next": next_link}}


class FakeRepresent:
    """Serves representative sets and representatives with real
    limit/offset pagination. `max_page` caps the page size (to exercise
    multi-page downloads)."""

    def __init__(self, sets, reps, max_page=1000):
        self.sets = list(sets)
        self.reps = list(reps)
        self.max_page = max_page
        self.bulk_total_override = None

    def _capped(self, url):
        query = urllib.parse.parse_qs(urllib.parse.urlsplit(url).query)
        limit = min(int(query.get("limit", ["20"])[0]), self.max_page)
        offset = int(query.get("offset", ["0"])[0])
        base = url.split("?", 1)[0]
        return f"{base}?limit={limit}&offset={offset}"

    def handle(self, url):
        path = urllib.parse.urlsplit(url).path
        if path == "/representative-sets/":
            return json_response(_page(self.sets, self._capped(url)))
        if path == "/representatives/":
            payload = _page(self.reps, self._capped(url))
            if self.bulk_total_override is not None:
                payload["meta"]["total_count"] = self.bulk_total_override
            return json_response(payload)
        for rep_set in self.sets:
            if path == rep_set["related"]["representatives_url"]:
                members = [r for r in self.reps if r["related"]["representative_set_url"] == rep_set["url"]]
                return json_response(_page(members, self._capped(url)))
        return http_error(url, 404)

    def install(self, web):
        web.route(REPRESENT, self.handle)
        return self


def sample_represent(**kwargs):
    data = load_fixture("represent_sample.json")
    return FakeRepresent(data["sets"], data["representatives"], **kwargs)


def profile_xml(committees=()):
    """A House of Commons member profile XML, shaped like the live one in
    fixtures/ourcommons_profile.xml, with the given (name, role, to_date) committees."""
    roles = "".join(
        "<CommitteeMemberRole><ParliamentNumber>45</ParliamentNumber><SessionNumber>1</SessionNumber>"
        f"<AffiliationRoleName>{escape(role)}</AffiliationRoleName><CommitteeName>{escape(name)}</CommitteeName>"
        "<FromDateTime>2025-06-13T12:03:00</FromDateTime>"
        + (f"<ToDateTime>{to_date}</ToDateTime>" if to_date else '<ToDateTime xsi:nil="true" />')
        + "</CommitteeMemberRole>"
        for name, role, to_date in committees
    )
    return (
        '<Profile xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance" '
        'xmlns:xsd="http://www.w3.org/2001/XMLSchema"><MemberOfParliamentRole><PersonId>1</PersonId>'
        "</MemberOfParliamentRole><CaucusMemberRoles /><ParliamentaryPositionRoles />"
        f"<CommitteeMemberRoles>{roles}</CommitteeMemberRoles></Profile>"
    ).encode("utf-8")


# Committees served for each MP in the sample (keyed by House of Commons person ID).
SAMPLE_COMMITTEES = {
    "111067": [("Finance", "Member", None)],
    "88567": [("Procedure and House Affairs", "Vice-Chair", None), ("Finance", "Member", "2025-09-01T00:00:00")],
    "105630": [("Public Accounts", "Chair", None)],
    "123418": [("Science and Research", "Member", None)],
    "88538": [],
}


class FakeFederalSources:
    """OpenParliament + House of Commons XML, from the fixtures, plus any
    roles pages given as {House of Commons person ID: HTML}."""

    def __init__(self, committees=None, details_override=None, roles=None):
        self.data = load_fixture("openparliament_sample.json")
        self.committees = SAMPLE_COMMITTEES if committees is None else committees
        self.details = dict(self.data["details"], **(details_override or {}))
        self.roles = roles or {}

    def openparliament(self, url):
        path = urllib.parse.urlsplit(url).path
        if path == "/politicians/":
            return json_response({"objects": self.data["list"], "pagination": {"next_url": None}})
        if path in self.details:
            return json_response(self.details[path])
        return http_error(url, 404)

    def ourcommons(self, url):
        person_id = url.rsplit("(", 1)[-1].split(")", 1)[0]
        if url.endswith("/roles"):
            page = self.roles.get(person_id)
            return FakeResponse(page) if page is not None else http_error(url, 404)
        if person_id not in self.committees:
            return http_error(url, 404)
        return FakeResponse(profile_xml(self.committees[person_id]))

    def install(self, web):
        web.route(OPENPARLIAMENT, self.openparliament)
        web.route("https://www.ourcommons.ca/Members/", self.ourcommons)
        return self


class FakeSenate:
    """The Senate's List and Contact Information tabs, from the fixtures
    (or the HTML / exception given for either)."""

    def __init__(self, senators=None, contacts=None):
        self.tabs = {
            "senatorslist": load_fixture("senate_list.html") if senators is None else senators,
            "senatorscontactinformation": load_fixture("senate_contacts.html") if contacts is None else contacts,
        }

    def handle(self, url):
        tab = urllib.parse.parse_qs(urllib.parse.urlsplit(url).query).get("displayFor", [""])[0]
        page = self.tabs.get(tab)
        if page is None:
            return http_error(url, 404)
        return page if isinstance(page, BaseException) else FakeResponse(page)

    def install(self, web):
        web.route(SENATE, self.handle)
        return self


def senate_list_without(*names):
    """The List fixture minus the rows for these senators ("Last, First" as listed)."""
    page = load_fixture("senate_list.html")
    for name in names:
        start = page.index(f">{name}</a>")
        page = page[:page.rindex("<tr>", 0, start)] + page[page.index("</tr>", start) + len("</tr>"):]
    return page


def nunavut_member_page(name, constituency, legislative=None, offices=(), duties=(), honourable=False,
                        biography="was elected in the general election held on October 27, 2025."):
    """An MLA's page in the other layout the parser reads: plain paragraphs
    with each section's details on one line, labels inline -- the way the
    live page's text reads when copied. `legislative` and each of `offices`
    are the text of one paragraph."""
    title = f"Honourable {name}" if honourable else name
    parts = [f"<h1>{escape(title)}</h1>", f"<p>Constituency {escape(constituency)}</p>"]
    if duties:
        parts.append("<h3>Duties</h3>" + "".join(f"<p>{escape(d)}</p>" for d in duties))
    if legislative:
        parts.append(f"<h3>Legislative Office</h3><p>{legislative}</p>")
    if offices:
        parts.append("<h3>Constituency Office</h3>" + "".join(f"<p>{o}</p>" for o in offices))
    parts.append(f"<p>{escape(name)} {escape(biography)}</p>")
    return (f"<!DOCTYPE html><html><head><title>{escape(title)} | Nunavut Legislative Assembly</title></head>"
            f"<body><nav><a href=\"/members/mla\">Members</a></nav><main>{''.join(parts)}</main></body></html>")


# The members in fixtures/nunavut_members.html other than John Main (whose
# page is fixtures/nunavut_member.html), by node number, plus a page that
# isn't a member's. Mike's and Pedersen's details are as published in
# September 2026; Joanasie's constituency office is from Open North's
# older copy of his page, and his page has no Legislative Office section.
NUNAVUT_PAGES = {
    "2669": nunavut_member_page(
        "David Joanasie", "South Baffin", honourable=True, duties=["Speaker of the Legislative Assembly"],
        offices=["P.O. Box 156 Cape Dorset, NU X0A 0C0 Phone: (867) 897-8753 "
                 "Email: <a href=\"mailto:davidjoanasie@southbaffinmla.ca\">Email</a>"]),
    "6490": nunavut_member_page(
        "Johnny Mike", "Pangnirtung",
        legislative="Phone: (867) 975-5135 Fax: (867) 975-5124 Email: jmike@assembly.nu.ca",
        offices=["P.O. Box 36 Pangnirtung, NU X0A 0R0 Phone: (867) 473-8220 Fax: (867) 473-8227 "
                 "Email: JohnnyMike@pangnirtungmla.ca"]),
    "9695": nunavut_member_page(
        "Fred Pedersen", "Cambridge Bay",
        legislative="Phone: (867) 975-5127 Fax: (867) 975-5108 Email: fpedersen@assembly.nu.ca",
        offices=["P.O. Box 1363 Cambridge Bay, NU X0B 0C0 Phone: (867) 983-3777 Fax: (867) 983-3778 "
                 "Email: fredpedersen@cambridgebaymla.ca"]),
    "41": "<html><head><title>Visitor Information | Nunavut Legislative Assembly</title></head>"
          "<body><h1>Visitor Information</h1><p>Tours of the Legislative Assembly run daily.</p></body></html>",
}


class FakeNunavut:
    """assembly.nu.ca: the members page and members' pages from the
    fixtures. `pages` ({node number: HTML or exception}) adds or replaces
    pages; `listing` replaces the members page."""

    def __init__(self, pages=None, listing=None):
        self.listing = load_fixture("nunavut_members.html") if listing is None else listing
        self.pages = dict(NUNAVUT_PAGES, **{"6485": load_fixture("nunavut_member.html")})
        self.pages.update(pages or {})

    def handle(self, url):
        path = urllib.parse.urlsplit(url).path.rstrip("/")
        if path == "/members/mla":
            page = self.listing
        else:
            page = self.pages.get(path[len("/node/"):]) if path.startswith("/node/") else None
        if page is None:
            return http_error(url, 404)
        return page if isinstance(page, BaseException) else FakeResponse(page)

    def install(self, web):
        web.route(NUNAVUT, self.handle)
        return self


class ScraperTestCase(unittest.TestCase):
    """Every test gets its own output folder and a quiet stdout
    (self.output collects what the scripts printed)."""

    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="politicians-test-")
        self.addCleanup(shutil.rmtree, self.tmp, ignore_errors=True)
        env = mock.patch.dict(os.environ, {rc.OUTPUT_DIR_ENV_VAR: self.tmp})
        env.start()
        self.addCleanup(env.stop)
        self.output = io.StringIO()
        redirect = contextlib.redirect_stdout(self.output)
        redirect.__enter__()
        self.addCleanup(redirect.__exit__, None, None, None)
        rc.clear_cache()

    @property
    def master_path(self):
        return os.path.join(self.tmp, "master.xlsx")

    def read_sheet(self, path, sheet):
        """A sheet as a list of {header: value} dicts."""
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True)
        try:
            rows = list(wb[sheet].iter_rows(values_only=True))
        finally:
            wb.close()
        header = rows[0]
        return [dict(zip(header, r)) for r in rows[1:]]

    def sheet_names(self, path):
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True)
        try:
            return list(wb.sheetnames)
        finally:
            wb.close()


def pages(total, page_size):
    return math.ceil(total / page_size)
