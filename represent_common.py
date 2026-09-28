"""
Shared helpers for scraping Canadian elected officials from the
Represent API (Open North): https://represent.opennorth.ca/

Represent aggregates elected officials at ALL levels of Canadian
government (federal, provincial/territorial, and municipal/local) in
one free API, so we don't need a separate scraper for every
legislature's website. At the time of writing it publishes 121
"representative sets" -- the House of Commons, 12 of the 13
provincial/territorial legislatures (there is no Nunavut set) and 108
municipal councils -- about 3,800 people in total.

Free tier limit: 60 requests/minute; past that Represent answers with
HTTP 503. Every request in this project goes through http_get() below,
which throttles per host and retries transient failures, so no loop
anywhere has to remember to sleep.

Everything is standard library except the Excel writer, which uses
openpyxl (install with: pip install openpyxl).
"""
import csv
import email.utils
import http.client
import json
import os
import re
import shutil
import sys
import tempfile
import threading
import time
import unicodedata
import urllib.error
import urllib.parse
import urllib.request
from collections import Counter, namedtuple
from datetime import datetime, timezone

BASE_URL = "https://represent.opennorth.ca"
USER_AGENT = "ca-politician-scraper/2.0 (personal research script)"


# --- Low-level HTTP -------------------------------------------------------
#
# Minimum gap between two requests to the same host. Represent allows 60
# requests/minute, so 1.1s keeps us at ~54/min no matter which loop is
# asking. (The throttle used to live in the pagination loops, which only
# slept *between pages of one council* -- a local scrape then fired its
# one-page councils back-to-back and ran straight into the rate limit.)
HOST_MIN_INTERVAL_SECONDS = {
    "represent.opennorth.ca": 1.1,
    "api.openparliament.ca": 1.0,
    "www.ourcommons.ca": 0.5,
}
DEFAULT_MIN_INTERVAL_SECONDS = 1.0

# Worth retrying: timeouts, rate limiting and server-side hiccups.
# Anything else (404, 400, ...) is permanent, so retrying only wastes time.
RETRYABLE_STATUSES = {408, 425, 429, 500, 502, 503, 504}
MAX_ATTEMPTS = 5          # for retryable HTTP statuses (backoff 5s..40s)
NETWORK_MAX_ATTEMPTS = 3  # for connection-level failures (backoff 2s, 4s)
MAX_RETRY_AFTER_SECONDS = 120

# Loops that make one request per item (per council, per MP) give up on
# the rest after this many consecutive failures: the site is down or
# unreachable, and trying hundreds more items would take hours.
MAX_CONSECUTIVE_FAILURES = 3

_throttle_lock = threading.Lock()
_next_request_at = {}


class FetchError(RuntimeError):
    """A URL couldn't be fetched: a permanent HTTP error, or a transient
    one that outlasted every retry. `status` is the HTTP status, if any."""

    def __init__(self, url, reason, status=None):
        super().__init__(f"Failed to fetch {url}: {reason}")
        self.url = url
        self.status = status


def _throttle(host):
    """Wait until `host` may receive another request. Thread-safe: each
    caller reserves the next free slot, so parallel workers hitting the
    same host still queue up at that host's minimum interval."""
    interval = HOST_MIN_INTERVAL_SECONDS.get(host, DEFAULT_MIN_INTERVAL_SECONDS)
    with _throttle_lock:
        now = time.monotonic()
        slot = max(now, _next_request_at.get(host, now))
        _next_request_at[host] = slot + interval
    if slot > now:
        time.sleep(slot - now)


def _retry_after_seconds(value):
    """Seconds from a Retry-After header (a number or an HTTP date), or None."""
    if not value:
        return None
    try:
        seconds = float(value)
    except ValueError:
        try:
            when = email.utils.parsedate_to_datetime(value)
        except (TypeError, ValueError):
            return None
        if when.tzinfo is None:
            when = when.replace(tzinfo=timezone.utc)
        seconds = (when - datetime.now(timezone.utc)).total_seconds()
    return min(max(seconds, 0.0), MAX_RETRY_AFTER_SECONDS)


def http_get(url, accept="application/json", headers=None, timeout=60):
    """
    GET `url` and return the raw response body (bytes).

    Every request in this project goes through here, so all of them get:
      - the per-host throttle above;
      - retries with exponential backoff for 408/429/5xx responses (5s,
        10s, 20s, 40s -- long enough to ride out a per-minute rate-limit
        window), honouring Retry-After;
      - a couple of quick retries for timeouts and dropped connections;
      - an immediate FetchError on permanent errors like 404.
    """
    host = (urllib.parse.urlsplit(url).hostname or "").lower()
    request_headers = {"User-Agent": USER_AGENT, "Accept": accept}
    request_headers.update(headers or {})

    http_failures = network_failures = 0
    while True:
        _throttle(host)
        try:
            req = urllib.request.Request(url, headers=request_headers)
            with urllib.request.urlopen(req, timeout=timeout) as resp:
                return resp.read()
        except urllib.error.HTTPError as e:
            if e.code not in RETRYABLE_STATUSES:
                raise FetchError(url, f"HTTP {e.code} {e.reason}", status=e.code) from None
            error, status = f"HTTP {e.code} {e.reason}", e.code
            http_failures += 1
            give_up = http_failures >= MAX_ATTEMPTS
            retry_after = _retry_after_seconds(e.headers.get("Retry-After") if e.headers else None)
            delay = retry_after if retry_after is not None else 2.5 * 2 ** http_failures
        except (OSError, http.client.HTTPException) as e:
            # OSError covers URLError (DNS failures, refused connections),
            # timeouts, and the ConnectionResetError / RemoteDisconnected
            # raised when a server drops the connection mid-response;
            # HTTPException covers IncompleteRead and other truncated or
            # malformed responses. (The old retry loop only caught URLError,
            # so a dropped connection ended the whole run.)
            error, status = f"{type(e).__name__}: {e}", None
            network_failures += 1
            give_up = network_failures >= NETWORK_MAX_ATTEMPTS
            delay = 2.0 * network_failures
        if give_up:
            raise FetchError(url, error, status=status)
        print(f"  ! request failed ({error}) -- retrying in {delay:.0f}s...")
        time.sleep(delay)


def http_get_json(url, headers=None):
    """http_get() + JSON decoding."""
    body = http_get(url, accept="application/json", headers=headers)
    try:
        return json.loads(body.decode("utf-8"))
    except (UnicodeDecodeError, ValueError) as e:
        raise FetchError(url, f"response wasn't valid JSON ({e})") from None


def _get_json(url):
    """Older name for http_get_json(), kept so existing callers still work."""
    return http_get_json(url)


def explain_fetch_error(error):
    """What the scrapers print when they can't reach a data source at all."""
    host = urllib.parse.urlsplit(getattr(error, "url", "") or "").hostname or "the data source"
    print(f"Couldn't reach {host} ({error}).\n"
          f"Check your internet connection or try again later -- nothing was changed.")


# --- Output location --------------------------------------------------------

OUTPUT_DIR_ENV_VAR = "CA_POLITICIANS_OUTPUT_DIR"


def _windows_desktop():
    """
    The Desktop folder Explorer actually shows, read from the registry.
    With OneDrive's "Known Folder Move" (on by default on many Windows
    10/11 machines) that's %OneDrive%\\Desktop rather than ~\\Desktop --
    writing to ~\\Desktop there creates a folder the user never sees.
    """
    try:
        import winreg
    except ImportError:
        return None
    try:
        with winreg.OpenKey(
            winreg.HKEY_CURRENT_USER,
            r"Software\Microsoft\Windows\CurrentVersion\Explorer\User Shell Folders",
        ) as key:
            value, _ = winreg.QueryValueEx(key, "Desktop")
    except OSError:
        return None
    path = os.path.expandvars(value)
    return path if os.path.isdir(path) else None


def get_output_root():
    """
    The top-level output folder: $CA_POLITICIANS_OUTPUT_DIR (which
    --output-dir sets) if given, otherwise a folder called "output" on the
    user's Desktop. Works on macOS, Windows (OneDrive-redirected Desktops
    included) and Linux.
    """
    override = os.environ.get(OUTPUT_DIR_ENV_VAR)
    if override:
        return os.path.abspath(os.path.expanduser(override))
    desktop = _windows_desktop() if sys.platform == "win32" else None
    if not desktop:
        desktop = os.path.join(os.path.expanduser("~"), "Desktop")
    return os.path.join(desktop, "output")


def get_output_dir(*subfolders):
    """Return (and create) a path inside the output folder."""
    out = os.path.join(get_output_root(), *subfolders)
    os.makedirs(out, exist_ok=True)
    return out


def get_master_path():
    """Path to the single persistent master workbook, shared across all scrapers."""
    return os.path.join(get_output_dir(), "master.xlsx")


# --- Command-line helpers ---------------------------------------------------

def add_common_args(parser):
    """Options every script accepts."""
    parser.add_argument(
        "--output-dir",
        metavar="DIR",
        help=f"write output here instead of ~/Desktop/output (same as setting {OUTPUT_DIR_ENV_VAR})",
    )
    return parser


def apply_common_args(args):
    """Act on the options from add_common_args(). Call once, right after parsing."""
    configure_console()
    if getattr(args, "output_dir", None):
        os.environ[OUTPUT_DIR_ENV_VAR] = os.path.abspath(os.path.expanduser(args.output_dir))


def configure_console():
    """Never crash on printing a name the console can't encode (e.g. output
    redirected to a log file on Windows, which uses a legacy code page)."""
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(errors="replace")
        except (AttributeError, ValueError):
            pass


# --- Classifying representative-sets by level of government ----------------

# Represent doesn't expose a clean "level" field on representative-sets,
# so we classify by the naming conventions Open North uses. Checked
# against all 121 sets Represent published at the time of writing:
# 1 federal, 12 provincial/territorial, 108 local, none left as "other".
FEDERAL_MARKERS = ("house of commons",)

PROVINCIAL_MARKERS = (
    "legislature",
    "legislative assembly",
    "national assembly",
    "assemblée nationale",
    "assemblée législative",
    "house of assembly",
)

MUNICIPAL_MARKERS = ("council", "conseil")


def classify_set(rep_set):
    """
    Classify a representative-set as 'federal', 'provincial_territorial',
    'local', or 'other'.
    """
    name = (rep_set.get("name") or "").lower()

    if any(marker in name for marker in FEDERAL_MARKERS):
        return "federal"

    if any(marker in name for marker in PROVINCIAL_MARKERS) and not any(
        marker in name for marker in MUNICIPAL_MARKERS
    ):
        return "provincial_territorial"

    if any(marker in name for marker in MUNICIPAL_MARKERS):
        return "local"

    return "other"


PROVINCES_AND_TERRITORIES = {
    "AB": "Alberta",
    "BC": "British Columbia",
    "MB": "Manitoba",
    "NB": "New Brunswick",
    "NL": "Newfoundland and Labrador",
    "NS": "Nova Scotia",
    "NT": "Northwest Territories",
    "NU": "Nunavut",
    "ON": "Ontario",
    "PE": "Prince Edward Island",
    "QC": "Quebec",
    "SK": "Saskatchewan",
    "YT": "Yukon",
}

# Statistics Canada province/territory codes: the first two digits of
# every federal electoral district (and census subdivision) code.
SGC_PROVINCE_CODES = {
    "10": "NL", "11": "PE", "12": "NS", "13": "NB", "24": "QC", "35": "ON", "46": "MB",
    "47": "SK", "48": "AB", "59": "BC", "60": "YT", "61": "NT", "62": "NU",
}


def province_for_set(rep_set):
    """
    Two-letter province/territory code for a representative-set, read from
    the scraper id at the end of its data_url ("…/ca_on_hamilton/" -> "ON",
    "…/ca_nt/" -> "NT"). Council names alone often can't tell you -- there
    are Richmond, Stratford and Waterloo councils in different provinces.
    Returns "" for the House of Commons ("…/ca/") or if unknown.
    """
    parts = [p for p in (rep_set.get("data_url") or "").split("/") if p]
    match = re.match(r"ca_([a-z]{2})(?:_|$)", parts[-1] if parts else "")
    if match and match.group(1).upper() in PROVINCES_AND_TERRITORIES:
        return match.group(1).upper()
    return ""


def province_for_representative(rep):
    """
    Province/territory for one representative when their set doesn't pin
    it down (the House of Commons spans the country): each MP's
    boundary_url ends in their riding's federal electoral district code,
    whose first two digits are the province (…/48007/ -> 48 -> "AB").
    """
    boundary_url = (rep.get("related") or {}).get("boundary_url") or ""
    match = re.search(r"/(\d{5})/?$", boundary_url)
    return SGC_PROVINCE_CODES.get(match.group(1)[:2], "") if match else ""


# --- Fetching data from Represent -------------------------------------------

# With this many sets or more, one bulk download of every representative
# (about 4 requests for ~3,800 people -- the method Open North recommends)
# beats one request per set.
BULK_FETCH_MIN_SETS = 6

_cache = {}


def clear_cache():
    """Forget API responses memoized during this run (mainly for tests)."""
    _cache.clear()


def _fetch_paginated(url):
    """Follow a Represent listing's meta.next links. Returns (objects, total_count)."""
    objects = []
    total = None
    visited = set()
    while url and url not in visited:
        visited.add(url)
        payload = http_get_json(url)
        objects.extend(payload.get("objects") or [])
        meta = payload.get("meta") or {}
        if meta.get("total_count") is not None:
            total = meta["total_count"]
        next_link = meta.get("next")
        url = urllib.parse.urljoin(BASE_URL, next_link) if next_link else None
    return objects, total


def fetch_all_representative_sets():
    """Return every representative-set object (one per legislature/council).
    Memoized, so scrape_all.py only asks once per run."""
    if "sets" not in _cache:
        print("Fetching the list of representative sets...")
        _cache["sets"], _ = _fetch_paginated(f"{BASE_URL}/representative-sets/?limit=1000")
    return list(_cache["sets"])


def fetch_representatives_by_url(representatives_url, set_name=None):
    """Return every representative belonging to one representative-set."""
    reps, _ = _fetch_paginated(f"{BASE_URL}{representatives_url}?limit=1000")
    print(f"  -> {set_name or representatives_url}: {len(reps)} representative(s)")
    return reps


def fetch_all_representatives():
    """
    Every representative in Represent, from the bulk endpoint
    (/representatives/?limit=1000). Memoized per run. Raises FetchError if
    the download fails or comes back short of the API's own total_count
    (e.g. because the data changed mid-download).
    """
    if "all_reps" not in _cache:
        print("Downloading all representatives in bulk...")
        url = f"{BASE_URL}/representatives/?limit=1000"
        reps, total = _fetch_paginated(url)
        if total is not None and len(reps) != total:
            raise FetchError(url, f"bulk download returned {len(reps)} of {total} representatives")
        print(f"  -> {len(reps)} representative(s)")
        _cache["all_reps"] = reps
    return list(_cache["all_reps"])


FetchResult = namedtuple("FetchResult", "sets rows fetched_sets failed_sets")


def fetch_rows_for_sets(rep_sets, category=None):
    """
    Fetch and flatten every representative in `rep_sets`; returns a
    FetchResult(sets, rows, fetched_sets, failed_sets), the last two being
    lists of set names.

    Many sets -> one bulk download; a few -> one request per set. A set
    that can't be downloaded is skipped and listed in `failed_sets`
    instead of aborting the run, and upsert_master() never marks people
    in a failed set as "removed".

    `category` is stamped on every row; None classifies each set.
    """
    reps_by_url = {}
    if rep_sets and (len(rep_sets) >= BULK_FETCH_MIN_SETS or "all_reps" in _cache):
        try:
            grouped = {}
            for rep in fetch_all_representatives():
                set_url = (rep.get("related") or {}).get("representative_set_url") or ""
                grouped.setdefault(set_url, []).append(rep)
            reps_by_url = {s["url"]: grouped.get(s["url"], []) for s in rep_sets}
        except FetchError as e:
            print(f"  ! bulk download failed ({e}) -- falling back to one request per set...")

    failed = []
    consecutive_failures = 0
    pending = [s for s in rep_sets if s["url"] not in reps_by_url]
    for i, rep_set in enumerate(pending, start=1):
        if consecutive_failures >= MAX_CONSECUTIVE_FAILURES:
            print(f"  ! Represent seems to be unreachable -- skipping the remaining {len(pending) - i + 1} set(s).")
            failed.extend(s["name"] for s in pending[i - 1:])
            break
        if len(pending) > 1:
            print(f"[{i}/{len(pending)}] {rep_set['name']}")
        try:
            reps_by_url[rep_set["url"]] = fetch_representatives_by_url(
                rep_set["related"]["representatives_url"], rep_set["name"]
            )
            consecutive_failures = 0
        except FetchError as e:
            print(f"  ! skipping {rep_set['name']}: {e}")
            failed.append(rep_set["name"])
            consecutive_failures += 1

    rows = []
    fetched = []
    for rep_set in rep_sets:
        reps = reps_by_url.get(rep_set["url"])
        if reps is None:
            continue
        fetched.append(rep_set["name"])
        set_category = category or classify_set(rep_set)
        province = province_for_set(rep_set)
        rows.extend(flatten_representative(r, set_category, province=province) for r in reps)
    return FetchResult(list(rep_sets), rows, fetched, failed)


def fetch_category(category, set_filter=None):
    """fetch_rows_for_sets() for one level of government ("federal",
    "provincial_territorial" or "local"), optionally narrowed by
    `set_filter(rep_set) -> bool`."""
    sets = [s for s in fetch_all_representative_sets() if classify_set(s) == category]
    if set_filter:
        sets = [s for s in sets if set_filter(s)]
    return fetch_rows_for_sets(sets, category)


def normalize_name(name):
    """
    Lowercase, strip accents/punctuation, and collapse whitespace so names
    from different sources (Represent vs. OpenParliament vs. Elections
    Canada) can be compared. Not perfect (doesn't handle nicknames), but
    good enough for matching "Marc Garneau" == "Marc Garneau".
    """
    if not name:
        return ""
    text = unicodedata.normalize("NFKD", str(name))
    text = "".join(c for c in text if not unicodedata.combining(c))
    text = text.lower()
    text = re.sub(r"[^a-z0-9\s]", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def slug_from_set(rep_set):
    """Derive a filesystem-friendly slug from a representative-set's url."""
    url = rep_set.get("url", "")
    parts = [p for p in url.split("/") if p]
    if parts:
        return parts[-1]
    return (rep_set.get("name") or "unknown").lower().replace(" ", "-")


# --- Flattening + writing output --------------------------------------------

CSV_FIELDS = [
    "category",
    "representative_set",
    "name",
    "first_name",
    "last_name",
    "elected_office",
    "district_name",
    "party_name",
    "email",
    "phone",
    "postal_address",
    "url",
    "personal_url",
    "photo_url",
    "gender",
    "source_url",
    # Added after v4 -- appended so column positions above never move.
    "province_territory",
    "constituency_phone",
    "constituency_address",
    "roles",
    "preferred_languages",
    "extra_info",
]

# Federal-only columns filled in by openparliament_enrich.py.
FEDERAL_ENRICHMENT_FIELDS = ["first_elected", "committees", "openparliament_url"]
FEDERAL_FIELDS = CSV_FIELDS + FEDERAL_ENRICHMENT_FIELDS

# Enrichment code lists the fields it successfully refreshed for a row
# under this key, so upsert_master() can tell "looked it up and it's now
# blank" (record it) from "wasn't looked up this run" (keep the old value).
# Keys starting with "_" are never written to any output file.
FRESH_FIELDS_KEY = "_fresh_fields"


def _value(obj, key):
    value = obj.get(key)
    return "" if value is None else value


def _flat_postal(postal):
    return (postal or "").replace("\n", ", ")


def _join(value):
    if isinstance(value, (list, tuple)):
        return "; ".join(str(v).strip() for v in value if v is not None and str(v).strip())
    return "" if value is None else str(value)


def _extra_info(extra, skip):
    """Every other `extra` key as "key: value; key: value" text, so nothing
    Represent publishes (e.g. Ontario's constituency_email) gets dropped."""
    parts = []
    for key in sorted(extra):
        if key in skip:
            continue
        value = extra[key]
        if isinstance(value, dict):
            value = json.dumps(value, ensure_ascii=False, sort_keys=True)
        elif isinstance(value, (list, tuple)):
            value = ", ".join(str(v) for v in value)
        if value not in (None, ""):
            parts.append(f"{key}: {value}")
    return "; ".join(parts)


def flatten_representative(rep, category, province=None):
    """Turn one API representative object into a flat dict for CSV export."""
    offices = rep.get("offices") or []
    primary_office = next(
        (o for o in offices if o.get("type") == "legislature"),
        offices[0] if offices else {},
    )
    # Legislators usually list their legislature office (for MPs, "House of
    # Commons, Ottawa" -- identical for all 343 of them) *and* a local
    # constituency office, which is the one constituents actually contact.
    constituency_office = next(
        (o for o in offices if o.get("type") == "constituency" and (o.get("tel") or o.get("postal"))),
        {},
    )
    extra = rep.get("extra") if isinstance(rep.get("extra"), dict) else {}
    return {
        "category": category,
        "representative_set": _value(rep, "representative_set_name"),
        "name": _value(rep, "name"),
        "first_name": _value(rep, "first_name"),
        "last_name": _value(rep, "last_name"),
        "elected_office": _value(rep, "elected_office"),
        "district_name": _value(rep, "district_name"),
        "party_name": _value(rep, "party_name"),
        "email": _value(rep, "email"),
        "phone": _value(primary_office, "tel"),
        "postal_address": _flat_postal(primary_office.get("postal")),
        "url": _value(rep, "url"),
        "personal_url": _value(rep, "personal_url"),
        "photo_url": _value(rep, "photo_url"),
        "gender": _value(rep, "gender"),
        "source_url": _value(rep, "source_url"),
        "province_territory": province or province_for_representative(rep),
        "constituency_phone": _value(constituency_office, "tel"),
        "constituency_address": _flat_postal(constituency_office.get("postal")),
        "roles": _join(extra.get("roles")),
        "preferred_languages": _join(extra.get("preferred_languages")),
        "extra_info": _extra_info(extra, skip=("roles", "preferred_languages")),
    }


def header_for_field(field):
    """"postal_address" -> "Postal Address" (the header text in every workbook)."""
    return field.replace("_", " ").title()


def field_for_header(header):
    """Inverse of header_for_field(): "Postal Address" -> "postal_address"."""
    return re.sub(r"\s+", "_", str(header or "").strip().lower())


def _public(row):
    """Drop internal bookkeeping keys (leading underscore) before output."""
    return {k: v for k, v in row.items() if not str(k).startswith("_")}


def _interactive():
    try:
        return sys.stdin is not None and sys.stdin.isatty()
    except (AttributeError, ValueError):
        return False


def _temp_path_for(path):
    """A temp file next to `path` (same folder, so the final rename is atomic)."""
    folder = os.path.dirname(os.path.abspath(path))
    fd, tmp = tempfile.mkstemp(prefix=".tmp-", suffix=os.path.splitext(path)[1], dir=folder)
    os.close(fd)
    return tmp


def _remove_quietly(path):
    try:
        os.remove(path)
    except OSError:
        pass


def _replace_file(tmp_path, path):
    """
    Move a fully-written temp file over `path` in one atomic step, so an
    interrupted run can never leave a half-written file behind.

    If `path` is locked -- on Windows, a workbook or CSV that's open in
    Excel can't be replaced -- ask the user to close it and retry when
    someone's at the keyboard; otherwise save the new data next to it as
    "<name> (unsaved <timestamp>)<ext>" and say so, instead of crashing at
    the very end of a long run and losing it. Returns the path written.
    """
    attempts = 3 if _interactive() else 1
    for attempt in range(1, attempts + 1):
        try:
            os.replace(tmp_path, path)
            return path
        except PermissionError:
            if attempt == attempts:
                break
            try:
                input(f"\n  '{path}' is open in another program (probably Excel).\n"
                      f"  Close it, then press Enter to retry... ")
            except (EOFError, OSError):
                break
    stem, ext = os.path.splitext(path)
    fallback = f"{stem} (unsaved {datetime.now().strftime('%Y%m%d-%H%M%S')}){ext}"
    os.replace(tmp_path, fallback)
    print(f"  ! Couldn't overwrite {path} -- it's open in another program.\n"
          f"    This run's data was saved to {fallback} instead.")
    return fallback


def _write_text_file(path, write):
    tmp = _temp_path_for(path)
    try:
        write(tmp)
    except BaseException:
        _remove_quietly(tmp)
        raise
    return _replace_file(tmp, path)


def write_csv(path, rows, fields=None):
    """Write `fields` (default CSV_FIELDS) of each row to a CSV file. Other
    keys on the rows are ignored -- they used to crash the run."""
    fields = fields or CSV_FIELDS

    def write(tmp):
        with open(tmp, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
            writer.writeheader()
            for row in rows:
                writer.writerow({k: _value(row, k) for k in fields})

    final = _write_text_file(path, write)
    print(f"Wrote {len(rows)} row(s) -> {final}")
    return final


def write_json(path, rows):
    def write(tmp):
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump([_public(r) for r in rows], f, ensure_ascii=False, indent=2, default=str)

    final = _write_text_file(path, write)
    print(f"Wrote {len(rows)} row(s) -> {final}")
    return final


# openpyxl refuses (raises IllegalCharacterError on) control characters
# that scraped text occasionally contains; tab/newline/CR are fine.
_ILLEGAL_XLSX_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _sheet_title(name):
    """Excel sheet names: at most 31 characters, none of []:*?/\\ ."""
    title = re.sub(r"[\[\]:*?/\\]", "-", name or "Sheet").strip("'")
    return title[:31] or "Sheet"


def _write_rows_to_sheet(ws, fields, rows):
    """Bold header, frozen top row, autofilter, rough column widths -- the
    formatting every workbook in this project shares."""
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    header_font = Font(name="Arial", bold=True)
    body_font = Font(name="Arial")

    for col_idx, field in enumerate(fields, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header_for_field(field))
        cell.font = header_font

    widths = [len(field) for field in fields]
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, field in enumerate(fields, start=1):
            value = row.get(field, "")
            if value is None:
                value = ""
            elif isinstance(value, str):
                value = _ILLEGAL_XLSX_CHARS.sub("", value)
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if isinstance(value, str) and value.startswith("="):
                cell.data_type = "s"  # data, never a live formula
            cell.font = body_font
            if value != "":
                widths[col_idx - 1] = max(widths[col_idx - 1], len(str(value)))

    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(fields))}{len(rows) + 1}"
    for col_idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = min(width + 2, 50)


def _open_or_create_workbook(path):
    from openpyxl import Workbook, load_workbook

    if os.path.exists(path):
        wb = load_workbook(path)
        # drop openpyxl's default blank "Sheet" if it's still sitting
        # there empty from a previous run that only just created the file
        if (
            "Sheet" in wb.sheetnames
            and wb["Sheet"].max_row == 1
            and wb["Sheet"].max_column == 1
            and wb["Sheet"]["A1"].value is None
        ):
            del wb["Sheet"]
        return wb
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def _reorder_sheets(wb, tab_order):
    order = [name for name in tab_order if name in wb.sheetnames]
    order += [name for name in wb.sheetnames if name not in order]
    wb._sheets = [wb[name] for name in order]


def _backup_path(path):
    stem, ext = os.path.splitext(path)
    return f"{stem}.backup{ext}"


def _save_workbook(wb, path, backup=False):
    """Save atomically (temp file + rename); with backup=True, first copy
    the previous version to "<name>.backup.xlsx". Returns the path written."""
    tmp = _temp_path_for(path)
    try:
        wb.save(tmp)
    except BaseException:
        _remove_quietly(tmp)
        raise
    if backup and os.path.exists(path):
        try:
            shutil.copy2(path, _backup_path(path))
        except OSError as e:
            print(f"  ! couldn't back up {path} ({e}) -- continuing")
    return _replace_file(tmp, path)


def write_xlsx(path, rows, sheet_name="Representatives", fields=None):
    """Write rows (dicts keyed by `fields`, default CSV_FIELDS) to a formatted .xlsx file."""
    try:
        from openpyxl import Workbook
    except ImportError:
        raise RuntimeError(
            "openpyxl is required for Excel output. Install it with: pip install openpyxl"
        )

    wb = Workbook()
    ws = wb.active
    ws.title = _sheet_title(sheet_name)
    _write_rows_to_sheet(ws, fields or CSV_FIELDS, rows)
    final = _save_workbook(wb, path)
    print(f"Wrote {len(rows)} row(s) -> {final}")
    return final


# --- Master workbook (persistent, upserted across runs) --------------------
#
# Unlike the per-category files above (which are just a fresh snapshot each
# run), the master workbook at ~/Desktop/output/master.xlsx is meant to be
# updated in place every time you re-run a scraper:
#   - new people are appended
#   - people already in the master have their fields updated if anything
#     changed (email, party, riding, etc. -- including a field going blank)
#   - people who disappear from a re-scraped category are marked
#     status="removed" rather than deleted, so history isn't lost; if they
#     come back later, the same row is marked "reactivated"
#   - every add/update/removal is appended as a row on the shared
#     "Changelog" sheet
#
# Each level of government gets its own tab -- Federal, Provincial_
# Territorial, Local -- exactly like MASTER_CANADA_POLITICIAN_BUDGET.xlsx's
# tab layout. scrape_federal.py only ever touches the Federal tab,
# scrape_local.py only ever touches Local, etc., so the three scrapers can
# run in any order without clobbering each other's data. Changelog is the
# one shared/append-only sheet across all three.
#
# Identity: Represent has no stable person ID. Its `url` field is the
# person's page on the legislature's own website -- blank for many
# councillors, and it changes whenever a legislature redesigns its site
# (or a council starts publishing profile pages). Keying on it alone made
# each such change look like everyone being "removed" and re-"added", with
# their history reset. So each row gets a permanent `record_id` when it's
# first seen, and a fresh scrape is matched to existing rows by, in order:
# the same set + same URL, the same set + same name, then the same set +
# same seat (district and office) + same last name -- each only when it's
# unambiguous on both sides. (The last rule is what lets "Bill Dennis"
# becoming "William Dennis" stay one person; the rare flip side is that a
# *different* Dennis winning that same seat shows up as a name change.)
#
# Columns are read by their header text, not position, so workbooks written
# by older versions (with fewer columns) load correctly, and columns a new
# version adds are filled in silently rather than logged as ~3,800 changes.

TRACKING_FIELDS = [
    "first_seen",
    "last_seen",
    "status",     # "active" or "removed"
    "record_id",  # permanent per-person id, shown as "key" in the Changelog
]

MASTER_FIELDS = CSV_FIELDS + FEDERAL_ENRICHMENT_FIELDS + TRACKING_FIELDS

CHANGELOG_FIELDS = [
    "timestamp",
    "category",
    "key",
    "name",
    "change_type",  # "added", "updated", "removed", "reactivated"
    "field",
    "old_value",
    "new_value",
]

# category_scope value -> the tab it lives on in master.xlsx
CATEGORY_TO_MASTER_SHEET = {
    "federal": "Federal",
    "provincial_territorial": "Provincial_Territorial",
    "local": "Local",
}

MASTER_TAB_ORDER = ["Federal", "Provincial_Territorial", "Local", "Changelog"]

# Represent fields: a fresh scrape is authoritative for these, blanks included.
BASE_COMPARE_FIELDS = [f for f in CSV_FIELDS if f != "category"]


def _now_iso():
    return datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")


def legacy_master_key(row):
    """How rows were keyed before record_id existed: the legislature-site
    URL, else set|name. Rows from older workbooks get this as their
    record_id, so their existing Changelog entries stay linked."""
    return row.get("url") or f"{row.get('representative_set', '')}|{row.get('name', '')}"


def master_key(row):
    """The stable identifier for a person in master.xlsx."""
    return row.get("record_id") or legacy_master_key(row)


def _load_sheet(path, sheet_name, fields):
    """
    One sheet as a list of dicts keyed by `fields`, matching columns by
    header text rather than position, plus the set of fields the sheet
    actually had a column for. ([], set()) if the file/sheet doesn't exist.
    """
    if not os.path.exists(path):
        return [], set()

    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            return [], set()
        rows_iter = wb[sheet_name].iter_rows(values_only=True)
        columns = [field_for_header(h) for h in (next(rows_iter, None) or ())]
        present = {c for c in columns if c in fields}
        rows = []
        for values in rows_iter:
            if not values or all(v is None or v == "" for v in values):
                continue
            row = {field: "" for field in fields}
            for column, value in zip(columns, values):
                if column in present and value is not None:
                    row[column] = value
            rows.append(row)
        return rows, present
    finally:
        wb.close()


def read_master_tab(path, sheet_name):
    """
    Load one level-of-government tab of master.xlsx into a dict of
    {key: row}. Returns {} if the file or that tab doesn't exist yet.
    """
    rows, _ = _load_sheet(path, sheet_name, MASTER_FIELDS)
    return {master_key(row): row for row in rows}


def read_all_master_rows(path):
    """
    Load every level-of-government tab combined into one dict of
    {key: row} -- for things like lookup_politician.py's cached search,
    which needs to search across all categories at once. Returns {} if
    the file doesn't exist yet.
    """
    combined = {}
    for sheet_name in ("Federal", "Provincial_Territorial", "Local"):
        for key, row in read_master_tab(path, sheet_name).items():
            combined[f"{sheet_name}|{key}"] = row
    return combined


def read_master_changelog(path):
    """Load the shared Changelog sheet. Returns [] if it doesn't exist yet."""
    rows, _ = _load_sheet(path, "Changelog", CHANGELOG_FIELDS)
    return rows


def _text(value):
    return "" if value is None else str(value)


def _identity_keys(row):
    """Ways to recognise the same person across runs, most specific first
    (see the Identity note above). None where the row lacks the data."""
    set_name = normalize_name(_text(row.get("representative_set")))
    url = _text(row.get("url")).strip()
    name = normalize_name(_text(row.get("name")))
    last = normalize_name(_text(row.get("last_name"))) or (name.split()[-1] if name else "")
    district = normalize_name(_text(row.get("district_name")))
    office = normalize_name(_text(row.get("elected_office")))
    return [
        (set_name, url) if url else None,
        (set_name, name) if name else None,
        (set_name, district, office, last) if district and last else None,
    ]


def match_to_existing(existing_rows, new_rows):
    """
    For each new row, the index of the existing master row for the same
    person, or None. A key only counts when it's unique among the new
    rows and among the not-yet-matched existing ones (an active row wins
    over a removed one), so e.g. a council that gives every member the
    same URL falls through to name matching instead of pairing the wrong
    people.
    """
    matches = [None] * len(new_rows)
    taken = set()
    new_keys = [_identity_keys(r) for r in new_rows]
    old_keys = [_identity_keys(r) for r in existing_rows]
    for level in range(3):
        new_counts = Counter(k[level] for k in new_keys if k[level] is not None)
        index = {}
        for i, keys in enumerate(old_keys):
            if i not in taken and keys[level] is not None:
                index.setdefault(keys[level], []).append(i)
        for j, keys in enumerate(new_keys):
            key = keys[level]
            if matches[j] is not None or key is None or new_counts[key] != 1:
                continue
            candidates = [i for i in index.get(key, []) if i not in taken]
            if len(candidates) > 1:
                active = [i for i in candidates if existing_rows[i].get("status") != "removed"]
                if len(active) == 1:
                    candidates = active
            if len(candidates) == 1:
                matches[j] = candidates[0]
                taken.add(candidates[0])
    return matches


def _unique_record_id(base, taken):
    record_id, n = base, 2
    while record_id in taken:
        record_id, n = f"{base}#{n}", n + 1
    return record_id


def upsert_master(master_path, new_rows, category_scope, attempted_sets=None,
                  category_complete=None):
    """
    Merge `new_rows` (a freshly-scraped batch, all belonging to
    `category_scope`, e.g. "federal") into that category's own tab of the
    persistent master workbook, and return the changelog entries
    generated by this run.

    Only touches the one tab for `category_scope` -- the other two levels'
    tabs are loaded and re-saved untouched, so running e.g.
    scrape_federal.py alone never wipes out your provincial/local data.

    Removal detection:
      - `attempted_sets`: names of the representative sets actually
        downloaded this run. People in other sets are left alone -- e.g.
        with scrape_local.py --filter, or when a council failed to
        download -- unless `category_complete` is true, meaning the run
        covered the whole category and those sets no longer exist in
        Represent at all.
      - A set that downloaded fine but came back with nobody in it is
        treated as an upstream glitch: its people are NOT marked removed.
      - attempted_sets=None keeps the old behaviour: everyone not in this
        batch is marked removed.
    """
    sheet_name = CATEGORY_TO_MASTER_SHEET.get(category_scope, category_scope.title())
    existing, present_fields = _load_sheet(master_path, sheet_name, MASTER_FIELDS)
    changelog_rows = read_master_changelog(master_path)
    now = _now_iso()
    new_entries = []

    def log(change_type, row, name, field="", old_value="", new_value=""):
        new_entries.append({
            "timestamp": now,
            "category": category_scope,
            "key": row["record_id"],
            "name": name,
            "change_type": change_type,
            "field": field,
            "old_value": old_value,
            "new_value": new_value,
        })

    taken_ids = set()
    for row in existing:
        row["record_id"] = _unique_record_id(_text(row.get("record_id")) or legacy_master_key(row), taken_ids)
        taken_ids.add(row["record_id"])

    matches = match_to_existing(existing, new_rows)
    seen = set()

    for new_row, match in zip(new_rows, matches):
        if match is None:
            merged = {field: _value(new_row, field) for field in MASTER_FIELDS}
            merged["first_seen"] = now
            merged["last_seen"] = now
            merged["status"] = "active"
            merged["record_id"] = _unique_record_id(
                f"{merged['representative_set']}|{merged['name']}", taken_ids
            )
            taken_ids.add(merged["record_id"])
            existing.append(merged)
            seen.add(len(existing) - 1)
            log("added", merged, merged["name"])
            continue

        seen.add(match)
        old_row = existing[match]
        name = _text(new_row.get("name")) or _text(old_row.get("name"))
        fresh = new_row.get(FRESH_FIELDS_KEY) or ()
        for field in BASE_COMPARE_FIELDS + FEDERAL_ENRICHMENT_FIELDS:
            if field not in new_row:
                continue  # not part of this scrape at all
            new_val = _text(new_row.get(field))
            if field in FEDERAL_ENRICHMENT_FIELDS and field not in fresh and not new_val:
                continue  # enrichment didn't run / didn't reach this person
            old_val = _text(old_row.get(field))
            if old_val == new_val:
                continue
            old_row[field] = new_val
            if field in present_fields:
                log("updated", old_row, name, field, old_val, new_val)
            # else: a column this version added -- fill in without logging
        old_row["last_seen"] = now
        if old_row.get("status") == "removed":
            log("reactivated", old_row, name)
        old_row["status"] = "active"

    if not new_rows:
        print("  ! Nothing was scraped for this category -- not marking anyone removed.")
    else:
        attempted = set(attempted_sets) if attempted_sets is not None else None
        complete = attempted is None if category_complete is None else category_complete
        rows_per_set = Counter(_text(r.get("representative_set")) for r in new_rows)
        empty_sets = set()
        for i, row in enumerate(existing):
            if i in seen or row.get("status") == "removed":
                continue
            set_name = _text(row.get("representative_set"))
            if attempted is not None:
                if set_name in attempted and rows_per_set[set_name] == 0:
                    empty_sets.add(set_name)
                    continue
                if set_name not in attempted and not complete:
                    continue
            row["status"] = "removed"
            log("removed", row, _text(row.get("name")))
        for set_name in sorted(empty_sets):
            print(f"  ! {set_name} came back empty -- assuming an upstream glitch, "
                  f"not marking its members removed.")

    write_master_tab(master_path, sheet_name, existing, changelog_rows + new_entries)
    return new_entries


def write_master_tab(path, sheet_name, tab_rows, changelog_rows):
    """
    Write/replace exactly one level-of-government tab of master.xlsx
    (creating the file if it doesn't exist yet), and rewrite the shared
    Changelog sheet with the full (existing + new) changelog, newest
    first. Any other level-of-government tab already in the workbook is
    preserved as-is. The previous version is kept as master.backup.xlsx.
    """
    wb = _open_or_create_workbook(path)

    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)
    _write_rows_to_sheet(ws, MASTER_FIELDS, sorted(tab_rows, key=lambda r: _text(r.get("name"))))

    if "Changelog" in wb.sheetnames:
        del wb["Changelog"]
    ws2 = wb.create_sheet("Changelog")
    changelog_sorted = sorted(changelog_rows, key=lambda r: _text(r.get("timestamp")), reverse=True)
    _write_rows_to_sheet(ws2, CHANGELOG_FIELDS, changelog_sorted)

    _reorder_sheets(wb, MASTER_TAB_ORDER)
    final = _save_workbook(wb, path, backup=True)
    print(
        f"'{sheet_name}' tab updated -> {final} "
        f"({len(tab_rows)} row(s), {len(changelog_rows)} changelog entrie(s))"
    )
    return final


# --- Shared multi-tab budget workbook (federal / provincial / local) -------
#
# MASTER_CANADA_POLITICIAN_BUDGET.xlsx holds one tab per level of
# government (Federal, Provincial_Territorial, Local) plus a Notes tab.
# Three separate scripts (scrape_budget_federal.py,
# scrape_budget_provincial_territorial.py, scrape_budget_local.py) each
# own exactly one data tab. write_budget_tab() below rewrites only the
# tab it's given and upserts that category's row in Notes -- every other
# tab in the workbook is loaded and re-saved untouched, so the three
# scripts can run in any order, independently, without wiping out each
# other's data.

BUDGET_TAB_ORDER = ["Federal", "Provincial_Territorial", "Local", "Notes"]


def _upsert_notes_row(wb, category, note_text):
    from openpyxl.styles import Font

    if "Notes" not in wb.sheetnames:
        ws = wb.create_sheet("Notes")
        for col_idx, h in enumerate(["Category", "Scope / Notes", "Last Updated"], start=1):
            ws.cell(row=1, column=col_idx, value=h).font = Font(name="Arial", bold=True)
        ws.column_dimensions["A"].width = 22
        ws.column_dimensions["B"].width = 110
        ws.column_dimensions["C"].width = 20
        ws.freeze_panes = "A2"
    else:
        ws = wb["Notes"]

    target_row = None
    for row_idx in range(2, ws.max_row + 1):
        if ws.cell(row=row_idx, column=1).value == category:
            target_row = row_idx
            break
    if target_row is None:
        target_row = ws.max_row + 1

    ws.cell(row=target_row, column=1, value=category).font = Font(name="Arial")
    ws.cell(row=target_row, column=2, value=note_text).font = Font(name="Arial", italic=True)
    ws.cell(row=target_row, column=3, value=_now_iso()).font = Font(name="Arial")


def write_budget_tab(path, sheet_name, fields, rows, notes_category, notes_text,
                     sort_field="total_contributions"):
    """
    Write/replace exactly one tab of the shared MASTER_CANADA_POLITICIAN_
    BUDGET.xlsx workbook (creating the file if it doesn't exist yet), and
    upsert that category's row on the Notes tab. Any other tabs already
    in the workbook -- other categories' budget data, other categories'
    notes -- are preserved as-is.
    """
    wb = _open_or_create_workbook(path)

    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)

    def _sort_key(row):
        val = row.get(sort_field)
        return val if isinstance(val, (int, float)) else -1

    rows_sorted = sorted(rows, key=_sort_key, reverse=True) if sort_field else list(rows)
    _write_rows_to_sheet(ws, fields, rows_sorted)

    _upsert_notes_row(wb, notes_category, notes_text)
    _reorder_sheets(wb, BUDGET_TAB_ORDER)

    final = _save_workbook(wb, path)
    print(f"'{sheet_name}' tab updated -> {final} ({len(rows_sorted)} row(s))")
    return final
